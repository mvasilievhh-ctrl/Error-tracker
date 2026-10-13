"""
ИИ-агент погоды: каждые 5 минут опрашивает погоду в Москве и соседних областях,
сохраняет замеры в SQLite и выгружает их в Excel (один файл на месяц).

Если задан ANTHROPIC_API_KEY, Claude раз в час пишет короткую сводку: общая картина,
где теплее и холоднее, что изменилось за час, есть ли опасные явления.

Запуск:
    python agent.py            # работать постоянно, опрос каждые 5 минут
    python agent.py --once     # один опрос и выход (для cron / Планировщика задач)
"""
import argparse
import json
import logging
import os
import sqlite3
import sys
import time
from contextlib import closing
from datetime import datetime, timedelta, timezone
from pathlib import Path

import anthropic
import requests
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

# ---------------------------------------------------------------------------
# Config (можно переопределить переменными окружения)
# ---------------------------------------------------------------------------
BASE_DIR = Path(__file__).resolve().parent
CITIES_FILE = Path(os.environ.get("CITIES_FILE", BASE_DIR / "cities.json"))
DATA_DIR = Path(os.environ.get("WEATHER_DATA_DIR", BASE_DIR / "data"))
DB_PATH = DATA_DIR / "weather.db"

CHECK_INTERVAL_MINUTES = float(os.environ.get("CHECK_INTERVAL_MINUTES", "5"))
AI_MODEL = os.environ.get("AI_MODEL", "claude-opus-5-5")
AI_INTERVAL_MINUTES = float(os.environ.get("AI_INTERVAL_MINUTES", "60"))

# Москва живёт по UTC+3 круглый год — так не нужен пакет часовых поясов на Windows
MSK = timezone(timedelta(hours=3), "MSK")

OPEN_METEO_URL = "https://api.open-meteo.com/v1/forecast"
# Не больше 10 полей: иначе Open-Meteo считает один запрос за несколько
CURRENT_FIELDS = [
    "temperature_2m",
    "apparent_temperature",
    "relative_humidity_2m",
    "surface_pressure",
    "wind_speed_10m",
    "wind_gusts_10m",
    "wind_direction_10m",
    "cloud_cover",
    "precipitation",
    "weather_code",
]
HPA_TO_MMHG = 0.750062

# Модели, которые поддерживают серверный fallback при отказе (refusal)
FALLBACK_MODELS = {"claude-opus-5-5", "claude-opus-5", "claude-fable-5-1", "claude-sonnet-5-5"}

log = logging.getLogger("weather_agent")

# ---------------------------------------------------------------------------
# Справочники
# ---------------------------------------------------------------------------

# Коды погоды WMO, которые возвращает Open-Meteo
WEATHER_CODES = {
    0: "Ясно",
    1: "Преимущественно ясно",
    2: "Переменная облачность",
    3: "Пасмурно",
    45: "Туман",
    48: "Туман с изморозью",
    51: "Слабая морось",
    53: "Морось",
    55: "Сильная морось",
    56: "Слабая ледяная морось",
    57: "Ледяная морось",
    61: "Небольшой дождь",
    63: "Дождь",
    65: "Сильный дождь",
    66: "Слабый ледяной дождь",
    67: "Ледяной дождь",
    71: "Небольшой снег",
    73: "Снег",
    75: "Сильный снег",
    77: "Снежные зёрна",
    80: "Слабый ливень",
    81: "Ливень",
    82: "Сильный ливень",
    85: "Небольшой снегопад",
    86: "Сильный снегопад",
    95: "Гроза",
    96: "Гроза с небольшим градом",
    99: "Гроза с сильным градом",
}

WIND_POINTS = ["С", "СВ", "В", "ЮВ", "Ю", "ЮЗ", "З", "СЗ"]


def describe_weather(code):
    if code is None:
        return ""
    return WEATHER_CODES.get(int(code), f"Код {int(code)}")


def describe_wind(degrees):
    if degrees is None:
        return ""
    point = WIND_POINTS[round(degrees / 45) % 8]
    return f"{point} ({round(degrees)}°)"


def now_msk():
    return datetime.now(MSK).replace(microsecond=0, tzinfo=None)


def load_cities():
    with open(CITIES_FILE, encoding="utf-8") as f:
        cities = json.load(f)
    if not cities:
        raise ValueError(f"В {CITIES_FILE} нет ни одного города")
    return cities

# ---------------------------------------------------------------------------
# Погода (Open-Meteo: бесплатно, без ключа)
# ---------------------------------------------------------------------------

def make_http_session():
    session = requests.Session()
    retry = Retry(total=3, backoff_factor=2, status_forcelist=[429, 500, 502, 503, 504])
    session.mount("https://", HTTPAdapter(max_retries=retry))
    return session


def fetch_weather(session, cities, polled_at):
    """Один запрос сразу за все города. Возвращает список строк для базы."""
    params = {
        "latitude": ",".join(str(c["lat"]) for c in cities),
        "longitude": ",".join(str(c["lon"]) for c in cities),
        "current": ",".join(CURRENT_FIELDS),
        "wind_speed_unit": "ms",
        "timezone": "Europe/Moscow",
    }
    resp = session.get(OPEN_METEO_URL, params=params, timeout=30)
    resp.raise_for_status()
    data = resp.json()
    if isinstance(data, dict):  # для одной точки API возвращает объект, а не список
        data = [data]
    if len(data) != len(cities):
        raise ValueError(f"Ожидали {len(cities)} точек, Open-Meteo вернул {len(data)}")

    rows = []
    for city, item in zip(cities, data):
        cur = item["current"]
        pressure = cur.get("surface_pressure")
        rows.append({
            "polled_at": polled_at.isoformat(sep=" "),
            "observed_at": cur["time"].replace("T", " "),
            "city": city["city"],
            "region": city.get("region", ""),
            "temperature": cur.get("temperature_2m"),
            "feels_like": cur.get("apparent_temperature"),
            "humidity": cur.get("relative_humidity_2m"),
            "pressure_mmhg": round(pressure * HPA_TO_MMHG, 1) if pressure is not None else None,
            "wind_speed": cur.get("wind_speed_10m"),
            "wind_gusts": cur.get("wind_gusts_10m"),
            "wind_direction": cur.get("wind_direction_10m"),
            "cloud_cover": cur.get("cloud_cover"),
            "precipitation": cur.get("precipitation"),
            "weather_code": cur.get("weather_code"),
        })
    return rows

# ---------------------------------------------------------------------------
# Database (полная история; Excel собирается из неё)
# ---------------------------------------------------------------------------

OBSERVATION_FIELDS = [
    "polled_at", "observed_at", "city", "region", "temperature", "feels_like",
    "humidity", "pressure_mmhg", "wind_speed", "wind_gusts", "wind_direction",
    "cloud_cover", "precipitation", "weather_code",
]


def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    conn = get_db()
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS observations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            polled_at TEXT NOT NULL,
            observed_at TEXT,
            city TEXT NOT NULL,
            region TEXT,
            temperature REAL,
            feels_like REAL,
            humidity REAL,
            pressure_mmhg REAL,
            wind_speed REAL,
            wind_gusts REAL,
            wind_direction REAL,
            cloud_cover REAL,
            precipitation REAL,
            weather_code INTEGER
        )
        """
    )
    conn.execute("CREATE INDEX IF NOT EXISTS idx_obs_polled ON observations(polled_at)")
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS summaries (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            created_at TEXT NOT NULL,
            observed_at TEXT,
            model TEXT,
            text TEXT NOT NULL
        )
        """
    )
    conn.commit()
    conn.close()


def save_observations(rows):
    placeholders = ", ".join("?" for _ in OBSERVATION_FIELDS)
    conn = get_db()
    conn.executemany(
        f"INSERT INTO observations ({', '.join(OBSERVATION_FIELDS)}) VALUES ({placeholders})",
        [[r[f] for f in OBSERVATION_FIELDS] for r in rows],
    )
    conn.commit()
    conn.close()


def get_snapshot_between(start, end):
    """Последний опрос в промежутке [start, end] (или пустой список)."""
    conn = get_db()
    row = conn.execute(
        "SELECT MAX(polled_at) AS t FROM observations WHERE polled_at BETWEEN ? AND ?",
        (start.isoformat(sep=" "), end.isoformat(sep=" ")),
    ).fetchone()
    rows = []
    if row["t"]:
        rows = conn.execute(
            "SELECT * FROM observations WHERE polled_at = ? ORDER BY id", (row["t"],)
        ).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def get_last_summary():
    conn = get_db()
    row = conn.execute("SELECT * FROM summaries ORDER BY id DESC LIMIT 1").fetchone()
    conn.close()
    return dict(row) if row else None


def save_summary(created_at, observed_at, model, text):
    conn = get_db()
    conn.execute(
        "INSERT INTO summaries (created_at, observed_at, model, text) VALUES (?, ?, ?, ?)",
        (created_at.isoformat(sep=" "), observed_at, model, text),
    )
    conn.commit()
    conn.close()

# ---------------------------------------------------------------------------
# AI-сводка (Claude)
# ---------------------------------------------------------------------------

SYSTEM_PROMPT = """Ты — метеоаналитик. Тебе присылают текущие замеры погоды в Москве, \
Московской области и соседних областях, а иногда ещё и замеры примерно часовой давности.

Напиши по-русски короткую сводку в 4–6 предложений:
- общая картина по региону;
- где сейчас теплее и холоднее всего;
- что заметно изменилось за час (если есть прошлые замеры);
- опасные явления, если они есть: порывы ветра от 15 м/с, сильные осадки, гроза, \
риск гололёда (температура около 0 °C и осадки), сильный мороз или жара.

Опирайся только на присланные цифры, ничего не придумывай. Пиши простым текстом, \
без заголовков, списков и markdown."""


def format_snapshot(rows):
    lines = ["Город | Регион | Темп °C | Ощущ °C | Влажн % | Давл мм | Ветер м/с | Порывы м/с | Напр | Облачн % | Осадки мм | Погода"]
    for r in rows:
        lines.append(" | ".join(str(v) for v in [
            r["city"], r["region"], r["temperature"], r["feels_like"], r["humidity"],
            r["pressure_mmhg"], r["wind_speed"], r["wind_gusts"],
            describe_wind(r["wind_direction"]), r["cloud_cover"], r["precipitation"],
            describe_weather(r["weather_code"]),
        ]))
    return "\n".join(lines)


def make_ai_client():
    client = anthropic.Anthropic()
    # Ключ берётся из ANTHROPIC_API_KEY (или профиля `ant auth login`); без него — работаем без сводок
    if client.api_key is None and client.auth_token is None and client.credentials is None:
        return None
    return client


def should_summarize(rows, polled_at):
    last = get_last_summary()
    if last is None:
        return True
    if last["observed_at"] == rows[0]["observed_at"]:
        return False  # Open-Meteo ещё не обновил данные — анализировать нечего
    last_time = datetime.fromisoformat(last["created_at"])
    return polled_at - last_time >= timedelta(minutes=AI_INTERVAL_MINUTES) - timedelta(seconds=30)


def write_ai_summary(client, rows, polled_at):
    previous = get_snapshot_between(polled_at - timedelta(minutes=115), polled_at - timedelta(minutes=55))
    prompt = f"Текущие замеры (данные на {rows[0]['observed_at']} МСК):\n{format_snapshot(rows)}"
    if previous:
        prompt += (
            f"\n\nЗамеры примерно час назад (данные на {previous[0]['observed_at']} МСК):\n"
            f"{format_snapshot(previous)}"
        )

    request = dict(
        model=AI_MODEL,
        max_tokens=16000,
        output_config={"effort": "low"},
        system=SYSTEM_PROMPT,
        messages=[{"role": "user", "content": prompt}],
    )
    if AI_MODEL in FALLBACK_MODELS:
        # При отказе модели запрос автоматически повторяется на рекомендованной модели
        request.update(betas=["server-side-fallback-2026-07-01"], fallbacks="default")

    try:
        response = client.beta.messages.create(**request)
    except anthropic.AuthenticationError:
        log.error("ИИ-сводка: неверный ANTHROPIC_API_KEY")
        return
    except anthropic.RateLimitError:
        log.warning("ИИ-сводка: превышен лимит запросов, попробую в следующий раз")
        return
    except anthropic.APIStatusError as e:
        log.warning("ИИ-сводка: ошибка API %s: %s", e.status_code, e.message)
        return
    except anthropic.APIConnectionError:
        log.warning("ИИ-сводка: нет связи с API Anthropic, попробую в следующий раз")
        return

    if response.stop_reason == "refusal":
        log.warning("ИИ-сводка: модель отказалась отвечать")
        return
    text = "\n".join(b.text for b in response.content if b.type == "text").strip()
    if not text:
        log.warning("ИИ-сводка: пустой ответ (stop_reason=%s)", response.stop_reason)
        return

    save_summary(polled_at, rows[0]["observed_at"], response.model, text)
    log.info("ИИ-сводка: %s", text)

# ---------------------------------------------------------------------------
# Excel (пересобирается из базы на каждом цикле, один файл на месяц)
# ---------------------------------------------------------------------------

WEATHER_COLUMNS = [
    ("Время опроса", 19),
    ("Время данных", 17),
    ("Город", 16),
    ("Регион", 22),
    ("Температура, °C", 13),
    ("Ощущается, °C", 13),
    ("Влажность, %", 12),
    ("Давление, мм рт. ст.", 13),
    ("Ветер, м/с", 10),
    ("Порывы, м/с", 10),
    ("Направление ветра", 14),
    ("Облачность, %", 12),
    ("Осадки, мм", 10),
    ("Погода", 26),
]
SUMMARY_COLUMNS = [("Время", 19), ("Данные на", 17), ("Модель", 18), ("Сводка", 120)]

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
HEADER_ALIGN = Alignment(wrap_text=True, vertical="center")
WRAP_ALIGN = Alignment(wrap_text=True, vertical="top")


def excel_path(moment):
    return DATA_DIR / f"weather_{moment:%Y-%m}.xlsx"


def to_datetime(value):
    return datetime.fromisoformat(value) if value else None


def weather_row(r):
    return [
        to_datetime(r["polled_at"]),
        to_datetime(r["observed_at"]),
        r["city"],
        r["region"],
        r["temperature"],
        r["feels_like"],
        r["humidity"],
        r["pressure_mmhg"],
        r["wind_speed"],
        r["wind_gusts"],
        describe_wind(r["wind_direction"]),
        r["cloud_cover"],
        r["precipitation"],
        describe_weather(r["weather_code"]),
    ]


def add_sheet(wb, title, columns, rows, wrap_last=False):
    ws = wb.create_sheet(title)
    for i, (_, width) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"

    header = []
    for name, _ in columns:
        cell = WriteOnlyCell(ws, value=name)
        cell.font, cell.fill, cell.alignment = HEADER_FONT, HEADER_FILL, HEADER_ALIGN
        header.append(cell)
    ws.append(header)

    count = 0
    for row in rows:
        if wrap_last:
            last = WriteOnlyCell(ws, value=row[-1])
            last.alignment = WRAP_ALIGN
            row = row[:-1] + [last]
        ws.append(row)
        count += 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{count + 1}"


def export_excel(moment):
    """Собирает файл за месяц: «Сейчас», «Журнал», «Сводки ИИ»."""
    month = f"{moment:%Y-%m}"
    path = excel_path(moment)
    tmp_path = path.with_name(path.stem + ".tmp.xlsx")

    with closing(get_db()) as conn:
        latest = conn.execute(
            "SELECT * FROM observations WHERE polled_at = (SELECT MAX(polled_at) FROM observations) ORDER BY id"
        ).fetchall()
        journal = conn.execute(
            "SELECT * FROM observations WHERE substr(polled_at, 1, 7) = ? ORDER BY polled_at, id",
            (month,),
        )
        summaries = conn.execute(
            "SELECT * FROM summaries WHERE substr(created_at, 1, 7) = ? ORDER BY id DESC",
            (month,),
        ).fetchall()

        # write_only — потоковая запись: быстро и мало памяти даже на сотнях тысяч строк
        wb = Workbook(write_only=True)
        add_sheet(wb, "Сейчас", WEATHER_COLUMNS, (weather_row(r) for r in latest))
        add_sheet(wb, "Журнал", WEATHER_COLUMNS, (weather_row(r) for r in journal))
        add_sheet(
            wb, "Сводки ИИ", SUMMARY_COLUMNS,
            ([to_datetime(s["created_at"]), to_datetime(s["observed_at"]), s["model"], s["text"]]
             for s in summaries),
            wrap_last=True,
        )

        wb.save(tmp_path)

    try:
        os.replace(tmp_path, path)  # атомарно: файл никогда не остаётся «наполовину записанным»
    except PermissionError:
        tmp_path.unlink(missing_ok=True)
        log.warning("Файл %s открыт в Excel — обновлю его на следующем цикле (данные сохранены в базе)", path.name)
        return None
    return path

# ---------------------------------------------------------------------------
# Main loop
# ---------------------------------------------------------------------------

def run_cycle(session, ai_client, cities):
    polled_at = now_msk()
    rows = fetch_weather(session, cities, polled_at)
    save_observations(rows)

    moscow = rows[0]
    log.info(
        "Опрошено городов: %d (данные на %s). %s: %s °C, %s",
        len(rows), moscow["observed_at"], moscow["city"], moscow["temperature"],
        describe_weather(moscow["weather_code"]).lower(),
    )

    if ai_client and should_summarize(rows, polled_at):
        write_ai_summary(ai_client, rows, polled_at)

    path = export_excel(polled_at)
    if path:
        log.info("Excel обновлён: %s", path)


def setup_logging():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    fmt = logging.Formatter("%(asctime)s %(levelname)s %(message)s", "%Y-%m-%d %H:%M:%S")
    for handler in (logging.StreamHandler(sys.stdout), logging.FileHandler(DATA_DIR / "agent.log", encoding="utf-8")):
        handler.setFormatter(fmt)
        log.addHandler(handler)
    log.setLevel(logging.INFO)


def main():
    parser = argparse.ArgumentParser(description="ИИ-агент погоды по Москве и соседним областям")
    parser.add_argument("--once", action="store_true", help="сделать один опрос и выйти")
    parser.add_argument("--interval", type=float, default=CHECK_INTERVAL_MINUTES,
                        help="интервал опроса в минутах (по умолчанию 5)")
    parser.add_argument("--no-ai", action="store_true", help="не писать ИИ-сводки")
    args = parser.parse_args()

    setup_logging()
    init_db()
    cities = load_cities()
    session = make_http_session()

    ai_client = None if args.no_ai else make_ai_client()
    if ai_client:
        log.info("ИИ-сводки включены: модель %s, не чаще раза в %g мин", AI_MODEL, AI_INTERVAL_MINUTES)
    else:
        log.info("ИИ-сводки выключены (нет ANTHROPIC_API_KEY или указан --no-ai) — собираю только данные")

    log.info("Старт: %d городов, опрос каждые %g мин, данные в %s", len(cities), args.interval, DATA_DIR)

    interval = args.interval * 60
    next_run = time.monotonic()
    try:
        while True:
            # Сбой одного опроса не должен останавливать агента
            try:
                run_cycle(session, ai_client, cities)
            except requests.RequestException as e:
                log.error("Не удалось получить погоду (%s), попробую в следующий раз", type(e).__name__)
                if args.once:
                    sys.exit(1)
            except Exception:
                log.exception("Ошибка в цикле опроса, попробую в следующий раз")
                if args.once:
                    sys.exit(1)
            if args.once:
                break
            next_run += interval
            while next_run <= time.monotonic():  # цикл затянулся — пропущенные запуски не догоняем
                next_run += interval
            time.sleep(next_run - time.monotonic())
    except KeyboardInterrupt:
        log.info("Остановлено пользователем")


if __name__ == "__main__":
    main()
